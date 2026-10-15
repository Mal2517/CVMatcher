import io, ipaddress, socket
from urllib.parse import urlparse

import requests

MAX_BYTES = 5 * 1024 * 1024


def read_cv(filename: str, data: bytes) -> str:
    if len(data) > MAX_BYTES:
        raise ValueError("File is larger than 5 MB.")
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n".join((p.extract_text() or "") for p in pdf.pages)
        if len(text.strip()) < 50:  # scanned PDF -> OCR fallback
            text = _ocr_pdf(data)
        return text
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)
    if name.endswith((".png", ".jpg", ".jpeg", ".webp")):
        return _ocr_image(data)
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Use PDF, DOCX, PNG, JPG or TXT.")


def _ocr_image(data: bytes) -> str:
    import pytesseract
    from PIL import Image
    return pytesseract.image_to_string(Image.open(io.BytesIO(data)))


def _ocr_pdf(data: bytes) -> str:
    import pdfplumber, pytesseract
    out = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for p in pdf.pages[:4]:
            out.append(pytesseract.image_to_string(p.to_image(resolution=200).original))
    return "\n".join(out)


def _is_public(host: str) -> bool:
    try:
        for info in socket.getaddrinfo(host, None):
            ip = ipaddress.ip_address(info[4][0])
            if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
                return False
        return True
    except Exception:
        return False


def read_job_url(url: str) -> str:
    u = urlparse(url)
    if u.scheme not in ("http", "https") or not u.hostname or not _is_public(u.hostname):
        raise ValueError("Invalid job link.")
    import trafilatura
    r = requests.get(url, timeout=12, headers={"User-Agent": "Mozilla/5.0"})
    r.raise_for_status()
    text = trafilatura.extract(r.text) or ""
    if len(text) < 100:
        raise ValueError("Could not read that page (many job sites block bots). Paste the job text instead.")
    return text
